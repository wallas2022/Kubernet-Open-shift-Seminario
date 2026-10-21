from pathlib import Path
import json,re,os
from docx import Document
from docx.shared import Inches,Pt,RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT

root=Path(os.environ.get('SEMINARIO_ROOT',Path.cwd()))
content=json.loads((root/'presentacion/fase-06/fuentes/DOC-06-02-contenido.json').read_text(encoding='utf-8'))
script=['# Guion de la presentación del sistema de reservas por seis grupos','','DOC-06-02 · Treinta diapositivas · 20 de octubre de 2026','','La presentación conjunta sigue los seis grupos del documento rector y culmina con una demostración funcional. Las notas explican el caso propio y separan evidencia local de hipótesis pendientes. La guía DOC-07-02 desarrolla la preparación y el registro de la demostración.','']
for number,slide in enumerate(content,1):
    script += [f"## Diapositiva {number} {slide['title']}",'']
    if slide.get('group'):script += [f"Responsable: Grupo {slide['group']}.",'']
    script += [slide['notes'],'','Fuentes: '+'; '.join(slide['sources']), '']
script_path=root/'presentacion/fase-06/DOC-06-02-guion-seis-grupos.md'
script_path.write_text('\n'.join(script),encoding='utf-8')

def clean(value):return value.replace('**','').replace('`','')
def build(source,output,subtitle,summary):
    lines=source.read_text(encoding='utf-8').splitlines()
    doc=Document();section=doc.sections[0]
    section.page_width=Inches(8.5);section.page_height=Inches(11)
    section.top_margin=section.bottom_margin=Inches(.82)
    section.left_margin=section.right_margin=Inches(.92)
    for name in ['Normal','Title','Subtitle','Heading 1','Heading 2','Heading 3','Caption']:
        style=doc.styles[name];style.font.name='Times New Roman';style.font.size=Pt(12);style.font.color.rgb=RGBColor(0,0,0)
        rf=style.element.get_or_add_rPr().rFonts
        for attr in list(rf.attrib):
            if 'Theme' in attr:del rf.attrib[attr]
        for attr in ['ascii','hAnsi','eastAsia','cs']:rf.set(qn('w:'+attr),'Times New Roman')
        style.paragraph_format.space_after=Pt(6)
        if name=='Title' or name.startswith('Heading'):style.font.bold=True
    doc.styles['Normal'].paragraph_format.line_spacing=1.08
    for style in doc.styles:
        for border in list(style.element.iter(qn('w:pBdr'))):border.getparent().remove(border)
    doc.styles['Normal'].paragraph_format.widow_control=True
    doc.add_paragraph(lines[0].lstrip('# '),'Title');doc.add_paragraph(subtitle)
    doc.add_paragraph(summary)
    doc.add_paragraph('Índice','Heading 1')
    p=doc.add_paragraph();field=OxmlElement('w:fldSimple');field.set(qn('w:instr'),'TOC \\o "1-2" \\h \\z \\u')
    run=OxmlElement('w:r');t=OxmlElement('w:t');t.text='Actualizar el índice en Word con F9.';run.append(t);field.append(run);p._p.append(field)
    update=OxmlElement('w:updateFields');update.set(qn('w:val'),'true');doc.settings.element.append(update)
    doc.add_page_break()
    i=1
    while i<len(lines):
        line=lines[i].strip()
        if not line:i+=1;continue
        if line.startswith('|'):
            raw=[]
            while i<len(lines) and lines[i].strip().startswith('|'):
                row=[clean(cell.strip()) for cell in lines[i].strip().strip('|').split('|')]
                if not all(re.fullmatch(r'[-: ]+',cell) for cell in row):raw.append(row)
                i+=1
            table=doc.add_table(rows=0,cols=len(raw[0]));table.autofit=False
            widths=[1.55,2.15,2.94] if len(raw[0])==3 else [6.64/len(raw[0])]*len(raw[0])
            for column,width in zip(table.columns,widths):column.width=Inches(width)
            for index,row in enumerate(raw):
                cells=table.add_row().cells
                for cell,width,value in zip(cells,widths,row):
                    cell.width=Inches(width);cell.text=value;cell.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
                    props=cell._tc.get_or_add_tcPr();borders=OxmlElement('w:tcBorders')
                    for edge in ['top','left','bottom','right']:
                        element=OxmlElement('w:'+edge);element.set(qn('w:val'),'single');element.set(qn('w:sz'),'4');element.set(qn('w:color'),'D9D9D9');borders.append(element)
                    props.append(borders)
                    margins=OxmlElement('w:tcMar')
                    for edge in ['top','bottom','left','right']:
                        element=OxmlElement('w:'+edge);element.set(qn('w:w'),'90');element.set(qn('w:type'),'dxa');margins.append(element)
                    props.append(margins)
                    if index==0:
                        shading=OxmlElement('w:shd');shading.set(qn('w:fill'),'000000');props.append(shading)
                        for p in cell.paragraphs:
                            for r in p.runs:r.font.color.rgb=RGBColor(255,255,255);r.bold=True
                    for p in cell.paragraphs:p.paragraph_format.space_after=Pt(3)
                trprops=table.rows[-1]._tr.get_or_add_trPr();trprops.append(OxmlElement('w:cantSplit'))
                if index==0:trprops.append(OxmlElement('w:tblHeader'))
            doc.add_paragraph();continue
        if line.startswith('## '):doc.add_paragraph(line[3:],'Heading 1')
        elif line.startswith('### '):doc.add_paragraph(line[4:],'Heading 2')
        else:doc.add_paragraph(clean(line))
        i+=1
    footer=section.footer.paragraphs[0];footer.alignment=2
    footer.add_run('Seminario · ')
    page=OxmlElement('w:fldSimple');page.set(qn('w:instr'),'PAGE');footer._p.append(page)
    for r in footer.runs:r.font.name='Times New Roman';r.font.size=Pt(12);r.font.color.rgb=RGBColor(0,0,0)
    output.parent.mkdir(parents=True,exist_ok=True);doc.save(output);print(output)

build(root/'documentacion/fase-01/fuentes/DOC-01-02-investigacion-seis-grupos.md',root/'documentacion/fase-01/word/DOC-01-02-investigacion-seis-grupos.docx','Marco teórico aplicado al caso de éxito propio','La investigación se organiza en seis grupos y vincula sus hipótesis con el sistema de reservas construido, las evidencias locales y el experimento OpenShift pendiente.')
build(script_path,root/'presentacion/fase-06/DOC-06-02-guion-seis-grupos.docx','Guion y fuentes de treinta diapositivas','Un bloque por grupo y una demostración conjunta del sistema de reservas; las notas mantienen la procedencia y límites de cada afirmación.')
build(root/'documentacion/fase-07/fuentes/DOC-07-02-guia-seis-grupos.md',root/'documentacion/fase-07/word/DOC-07-02-guia-seis-grupos.docx','Organización exposición y demostración del sistema de reservas','Guía para coordinar 22 alumnos en seis grupos, preparar una demostración funcional y evaluar aportes aplicados al mismo caso.')
