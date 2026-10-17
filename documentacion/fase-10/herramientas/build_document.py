from pathlib import Path
import re, json
from docx import Document
from docx.shared import Pt, Inches, RGBColor, Cm
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from PIL import Image

ROOT = Path('/work')
BASE = ROOT / 'documentacion/fase-10'
OUT = BASE / 'word/DOC-10-00-analisis-diseno-frontend.docx'
FIGURES = BASE / 'revision/figuras'

doc = Document()
section = doc.sections[0]
section.page_width = Cm(21.59)
section.page_height = Cm(27.94)
section.top_margin = section.bottom_margin = Cm(2.2)
section.left_margin = section.right_margin = Cm(2.3)
for name in ['Normal','Title','Subtitle','Heading 1','Heading 2','Heading 3','Caption']:
    s = doc.styles[name]
    s.font.name = 'Times New Roman'
    s.font.size = Pt(12)
    s.font.color.rgb = RGBColor(0,0,0)
    rf=s.element.get_or_add_rPr().rFonts
    for key in list(rf.attrib):
        if 'Theme' in key:del rf.attrib[key]
    for key in ['ascii','hAnsi','eastAsia','cs']:rf.set(qn('w:'+key),'Times New Roman')
    s.paragraph_format.space_after = Pt(7)
    if name.startswith('Heading') or name == 'Title': s.font.bold = True
for style in doc.styles:
    for border in list(style.element.iter(qn('w:pBdr'))):border.getparent().remove(border)
doc.styles['Normal'].paragraph_format.line_spacing = 1.12
doc.styles['Normal'].paragraph_format.widow_control = True

doc.add_paragraph('Análisis y diseño del frontend de reservas', 'Title')
doc.add_paragraph('Seminario Kubernetes y Red Hat OpenShift')
doc.add_paragraph('DOC 10 00 | Fase 10 | Versión 1.0 | 17 de octubre de 2026')
doc.add_paragraph('Propuesta disponible para revisión')
doc.add_paragraph('Este documento reúne cinco etapas para construir la interfaz web de la aplicación de reservas. Define requisitos, casos de uso, pantallas, arquitectura y validación, reutilizando la API existente y su protección contra reservas duplicadas. Está dirigido al responsable del seminario y a quienes implementarán y revisarán el frontend.')
doc.add_paragraph('El frontend todavía no está implementado. Los diagramas describen el diseño propuesto, salvo el modelo de datos identificado como existente. Los casos de prueba son objetivos de verificación y no resultados ejecutados.')
doc.add_paragraph('La documentación editable por etapa se conserva en Markdown y los diez diagramas en Mermaid bajo la raíz del proyecto.')
doc.add_paragraph('Índice', 'Heading 1')
p = doc.add_paragraph()
field = OxmlElement('w:fldSimple')
field.set(qn('w:instr'), 'TOC \\o "1-2" \\h \\z \\u')
r = OxmlElement('w:r'); t = OxmlElement('w:t')
t.text = 'Actualizar este índice en Word con F9.'
r.append(t); field.append(r); p._p.append(field)
settings = doc.settings.element
upd=OxmlElement('w:updateFields');upd.set(qn('w:val'),'true');settings.append(upd)

figure_map = {
 'DOC-10-01': ['D10-01'], 'DOC-10-02':['D10-02'],
 'DOC-10-03':['D10-03','D10-04','D10-05'],
 'DOC-10-04':['D10-06','D10-07','D10-08','D10-09'],
 'DOC-10-05':['D10-10']
}
titles = {
 'D10-01':'Contexto del sistema', 'D10-02':'Mapa de casos de uso',
 'D10-03':'Navegación entre pantallas', 'D10-04':'Estados de la operación',
 'D10-05':'Distribución funcional de pantallas', 'D10-06':'Componentes propuestos',
 'D10-07':'Reserva y recuperación de respuesta perdida', 'D10-08':'Modelo de datos existente',
 'D10-09':'Despliegue propuesto', 'D10-10':'Dependencias de entrega'
}

def clean(text):
    return re.sub(r'`([^`]+)`',r'\1',text).replace('**','')

def table(lines):
    rows = [[clean(c.strip()) for c in line.strip().strip('|').split('|')] for line in lines]
    rows = [r for r in rows if not all(re.fullmatch(r'[:\- ]+', c or ' ') for c in r)]
    tab=doc.add_table(rows=0,cols=len(rows[0])); tab.autofit=False
    available=6.69
    if len(rows[0])==5:widths=[1.0,1.05,1.0,1.55,2.09]
    elif rows[0][0]=='Situación':widths=[1.8,2.8,2.09]
    elif rows[0][0]=='Riesgo':widths=[1.8,2.5,2.39]
    elif rows[0][0]=='Ruta del navegador':widths=[2.3,2.2,2.19]
    elif len(rows[0])==3 and rows[0][1]=='Pantalla':widths=[0.65,1.6,4.44]
    elif len(rows[0])==3 and rows[0][1]=='Escenario':widths=[0.65,3.15,2.89]
    elif len(rows[0])==3:widths=[0.65,3.65,2.39]
    elif len(rows[0])==2:widths=[2.2,4.49]
    else:widths=[available/len(rows[0])]*len(rows[0])
    for c,w in zip(tab.columns,widths): c.width=Inches(w)
    for i,row in enumerate(rows):
        cells=tab.add_row().cells
        for cell,w,value in zip(cells,widths,row):
            cell.width=Inches(w); cell.text=value
            tcpr=cell._tc.get_or_add_tcPr()
            borders=OxmlElement('w:tcBorders')
            for edge in ['top','left','bottom','right']:
                e=OxmlElement('w:'+edge); e.set(qn('w:val'),'single');e.set(qn('w:sz'),'4');e.set(qn('w:color'),'D9D9D9');borders.append(e)
            tcpr.append(borders)
            margins=OxmlElement('w:tcMar')
            for side in ['top','left','bottom','right']:
                e=OxmlElement('w:'+side);e.set(qn('w:w'),'85');e.set(qn('w:type'),'dxa');margins.append(e)
            tcpr.append(margins)
            valign=OxmlElement('w:vAlign');valign.set(qn('w:val'),'center');tcpr.append(valign)
            for p in cell.paragraphs:
                p.paragraph_format.space_after=Pt(3)
                p.paragraph_format.line_spacing=1.0
            if i==0:
                shd=OxmlElement('w:shd');shd.set(qn('w:fill'),'000000');tcpr.append(shd)
                for p in cell.paragraphs:
                    for run in p.runs:run.font.color.rgb=RGBColor(255,255,255);run.bold=True
        trpr=tab.rows[-1]._tr.get_or_add_trPr()
        cant=OxmlElement('w:cantSplit');trpr.append(cant)
        if i==0: trpr.append(OxmlElement('w:tblHeader'))
    doc.add_paragraph()

for path in sorted((BASE/'fuentes').glob('DOC-10-0[1-5]-*.md')):
    doc.add_page_break()
    lines=path.read_text(encoding='utf-8').splitlines()
    i=0
    while i<len(lines):
        line=lines[i].strip()
        if not line:i+=1;continue
        if line == '## Diagramas editables':break
        if line.startswith('```'):
            i+=1
            while i<len(lines) and not lines[i].strip().startswith('```'):i+=1
            i+=1
            continue
        if line.startswith('|'):
            data=[]
            while i<len(lines) and lines[i].strip().startswith('|'):data.append(lines[i]);i+=1
            table(data);continue
        if line.startswith('# '):doc.add_paragraph(clean(line[2:]),'Heading 1')
        elif line.startswith('## '):doc.add_paragraph(clean(line[3:]),'Heading 2')
        elif line.startswith('### '):doc.add_paragraph(clean(line[4:]),'Heading 3')
        else:doc.add_paragraph(clean(line))
        i+=1
    for fig in figure_map[path.name[:9]]:
        doc.add_page_break()
        doc.add_paragraph(fig+' '+titles[fig], 'Heading 2')
        image_path=next(FIGURES.glob(fig+'-*.png'))
        with Image.open(image_path) as im:w,h=im.size
        # Fit without exceeding the printable figure page.
        width=min(6.69, 7.65*w/h)
        picture=doc.add_picture(str(image_path),width=Inches(width))
        picture._inline.docPr.set('descr',titles[fig]+'. Interpretación y límites en la etapa correspondiente.')
        doc.add_paragraph('Figura '+fig+'. '+titles[fig]+'. Fuente editable: diagramas/fase-10/'+image_path.stem+'.mmd.', 'Caption')
        doc.add_paragraph('Modelo existente.' if fig=='D10-08' else 'Diseño propuesto. No representa una implementación ni una prueba ejecutada.')

footer=section.footer.paragraphs[0];footer.alignment=2
footer.add_run('DOC 10 00 | ')
fld=OxmlElement('w:fldSimple');fld.set(qn('w:instr'),'PAGE');footer._p.append(fld)
OUT.parent.mkdir(parents=True,exist_ok=True)
doc.save(OUT)
from zipfile import ZipFile
from lxml import etree
with ZipFile(OUT) as z:
    xml=etree.fromstring(z.read('word/document.xml'))
    ns={'w':'http://schemas.openxmlformats.org/wordprocessingml/2006/main',
        'a':'http://schemas.openxmlformats.org/drawingml/2006/main'}
    checks={
        'etapas_fuente':len(list((BASE/'fuentes').glob('DOC-10-0[1-5]-*.md'))),
        'figuras_word':len(xml.findall('.//a:blip',ns)),
        'indice_actualizable':bool(xml.xpath('.//w:fldSimple[contains(@w:instr,"TOC")]',namespaces=ns)),
        'encabezados_tabla_repetidos':len(xml.findall('.//w:tblHeader',ns)),
        'tablas':len(xml.findall('.//w:tbl',ns)),
        'alcance':'Comprobación documental; no ejecuta pruebas funcionales del frontend'
    }
    assert checks['etapas_fuente']==5 and checks['figuras_word']==10
    assert checks['indice_actualizable'] and checks['tablas']==checks['encabezados_tabla_repetidos']
    (BASE/'revision/estructura.json').write_text(json.dumps(checks,ensure_ascii=False,indent=2),encoding='utf-8')
print(OUT)
