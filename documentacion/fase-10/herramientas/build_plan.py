from pathlib import Path
from docx import Document
from docx.shared import Pt, Cm, Inches, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.image.image import Image
import re, argparse

parser=argparse.ArgumentParser()
parser.add_argument('--source',default='fuentes/DOC-10-06-plan-sprints-codificacion.md')
parser.add_argument('--output',default='word/DOC-10-06-plan-sprints-codificacion.docx')
parser.add_argument('--subtitle',default='Planificación de 19 sprints y aprobación de publicación')
parser.add_argument('--expected-sprints',type=int,default=19)
parser.add_argument('--figure')
parser.add_argument('--figure-max-height',type=float,default=8.5)
parser.add_argument('--figure-title',default='Frontera de integración de S01')
parser.add_argument('--figure-caption',default='Fuente editable: diagramas/fase-10/sprints/S01-base-frontend.mmd. La pantalla inicial no solicita credenciales ni realiza operaciones de negocio.')
parser.add_argument('--footer-label')
parser.add_argument('--include-code', action='store_true', help='Include command blocks in reproducible manuals')
args=parser.parse_args()

base=Path('/work/documentacion/fase-10')
source=base/args.source
out=base/args.output
doc=Document()
s=doc.sections[0]
s.page_width=Cm(21.59);s.page_height=Cm(27.94)
s.top_margin=s.bottom_margin=Cm(2.2)
s.left_margin=s.right_margin=Cm(2.3)
for name in ['Normal','Title','Heading 1','Heading 2','Heading 3']:
    st=doc.styles[name];st.font.name='Times New Roman';st.font.size=Pt(12);st.font.color.rgb=RGBColor(0,0,0)
    rf=st.element.get_or_add_rPr().rFonts
    for key in list(rf.attrib):
        if 'Theme' in key:del rf.attrib[key]
    for key in ['ascii','hAnsi','eastAsia','cs']:rf.set(qn('w:'+key),'Times New Roman')
    st.paragraph_format.space_after=Pt(7)
    if name!='Normal':st.font.bold=True
for st in doc.styles:
    for border in list(st.element.iter(qn('w:pBdr'))):border.getparent().remove(border)
doc.styles['Normal'].paragraph_format.line_spacing=1.1
doc.styles['Normal'].paragraph_format.widow_control=True

def clean(text):return re.sub(r'`([^`]+)`',r'\1',text)

def addtable(lines):
    rows=[[clean(c.strip()) for c in l.strip('|').split('|')] for l in lines]
    rows=[r for r in rows if not all(re.fullmatch(r'[-: ]+',c or ' ') for c in r)]
    t=doc.add_table(rows=0,cols=len(rows[0]));t.autofit=False
    widths=[.7,.9,4.2,.89] if len(rows[0])==4 else [1.6,2.6,2.49]
    for col,w in zip(t.columns,widths):col.width=Inches(w)
    for idx,row in enumerate(rows):
        cells=t.add_row().cells
        for cell,w,text in zip(cells,widths,row):
            cell.width=Inches(w);cell.text=text;pr=cell._tc.get_or_add_tcPr()
            border=OxmlElement('w:tcBorders')
            for edge in ['top','left','bottom','right']:
                e=OxmlElement('w:'+edge);e.set(qn('w:val'),'single');e.set(qn('w:sz'),'4');e.set(qn('w:color'),'D9D9D9');border.append(e)
            pr.append(border)
            if idx==0:
                fill=OxmlElement('w:shd');fill.set(qn('w:fill'),'000000');pr.append(fill)
                for p in cell.paragraphs:
                    for r in p.runs:r.font.color.rgb=RGBColor(255,255,255);r.bold=True
            for p in cell.paragraphs:p.paragraph_format.space_after=Pt(4)
        trpr=t.rows[-1]._tr.get_or_add_trPr();trpr.append(OxmlElement('w:cantSplit'))
        if idx==0:trpr.append(OxmlElement('w:tblHeader'))
    doc.add_paragraph()

lines=source.read_text(encoding='utf-8').splitlines()
doc.add_paragraph(lines[0][2:],'Title')
doc.add_paragraph(args.subtitle)
doc.add_paragraph('Índice','Heading 1')
p=doc.add_paragraph();f=OxmlElement('w:fldSimple');f.set(qn('w:instr'),'TOC \\o "1-2" \\h \\z \\u')
r=OxmlElement('w:r');t=OxmlElement('w:t');t.text='Actualizar el índice en Word con F9.';r.append(t);f.append(r);p._p.append(f)
doc.add_page_break()
i=1
while i<len(lines):
    line=lines[i].strip()
    if not line:i+=1;continue
    if line.startswith('```'):
        i+=1
        while i<len(lines) and not lines[i].strip().startswith('```'):
            if args.include_code:doc.add_paragraph(lines[i])
            i+=1
        i+=1
        continue
    if line.startswith('|'):
        data=[]
        while i<len(lines) and lines[i].strip().startswith('|'):data.append(lines[i].strip());i+=1
        addtable(data);continue
    if line.startswith('## '):doc.add_paragraph(line[3:],'Heading 1')
    elif line.startswith('### '):doc.add_paragraph(line[4:],'Heading 2')
    else:doc.add_paragraph(clean(line))
    i+=1
if args.figure:
    heading=doc.add_paragraph(args.figure_title,'Heading 1')
    heading.paragraph_format.page_break_before=True
    figure=Image.from_file(args.figure)
    width=min(6.69,args.figure_max_height*figure.px_width/figure.px_height)
    doc.add_picture(args.figure,width=Inches(width))
    doc.add_paragraph(args.figure_caption)
upd=OxmlElement('w:updateFields');upd.set(qn('w:val'),'true');doc.settings.element.append(upd)
footer=s.footer.paragraphs[0];footer.alignment=2;footer.add_run((args.footer_label or ('S01 RN03' if args.expected_sprints==0 else 'DOC 10 06'))+' | ')
field=OxmlElement('w:fldSimple');field.set(qn('w:instr'),'PAGE');footer._p.append(field)
out.parent.mkdir(parents=True,exist_ok=True)
doc.save(out)
if args.expected_sprints:
    assert len(re.findall(r'^### S\d\d ',source.read_text(encoding='utf-8'),re.M))==args.expected_sprints
print(out)
